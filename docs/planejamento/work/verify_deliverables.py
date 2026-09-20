import json,hashlib,re
from pathlib import Path
import openpyxl

tasks=json.loads(Path('work/backlog.json').read_text(encoding='utf-8'))
bykey={t['key']:t for t in tasks}
ids=[t['id'] for t in tasks]
assert len(ids)==len(set(ids))==127
assert ids==[f'BK-{i:03}' for i in range(len(ids))]
def ancestors(k):
    result=set()
    def run(x):
        for v in bykey[x]['deps']:
            if v not in result:result.add(v);run(v)
    run(k);return result
required=['usersscreen','parametersscreen','integrationsscreen','reconciliationscreen','receivablesscreen','payablesscreen','cashscreen','drescreen','trialbalancescreen','calculatorscreen','applicationsscreen','pricing','logisticsanalysis','inventoryscreen','productionscreen','purchasesscreen','projectionsscreen','projectionexport','pendingscreen','alertsscreen','executivescreen','extrascreens','approvedexports','cyclemetrics']
assert set(required)<=ancestors('readiness'),set(required)-ancestors('readiness')
assert set(['receivableingest','payableingest','logisticsingest'])<=ancestors('dreengine')
assert set(['restoreexercise','backupdesign','staginglocal'])<=ancestors('golive')
assert all(t['goal'] and t['action'] and len(t['steps'])>=3 and t['accept'] and t['impact'] for t in tasks)
md=Path('outputs/MAJUCAU_BACKLOG_DETALHADO_v0.md').read_text(encoding='utf-8')
assert re.findall(r'^### (BK-\d+) —',md,flags=re.M)==ids
assert not any(t['progress'] for t in tasks if t['kind']!='Planejamento')

p=Path('outputs/MAJUCAU_BACKLOG_E_EVOLUCAO_v0.xlsx')
w=openpyxl.load_workbook(p,data_only=False)
c=openpyxl.load_workbook(p,data_only=True)
assert w.sheetnames==['Evolucao','Tarefas','Decisoes','Cobertura']
assert [w['Tarefas'].cell(i+6,1).value for i in range(len(tasks))]==ids
assert c['Evolucao']['B6'].value==len(tasks)
assert c['Evolucao']['B7'].value==3
assert abs(c['Evolucao']['B8'].value-3/127)<1e-10
assert c['Evolucao']['B9'].value==0
assert c['Evolucao']['B10'].value==52
errors=[]
for s in c:
    for row in s:
        for cell in row:
            if cell.data_type=='e':errors.append((s.title,cell.coordinate,cell.value))
assert not errors,errors
assert all(w['Tarefas'].cell(i+6,7).data_type=='f' for i in range(len(tasks)))
assert len(w['Tarefas'].data_validations.dataValidation)>=2

manifest=json.loads(Path('work/source_extract/manifest.json').read_text(encoding='utf-8'))
root=Path(r'sources/private/material-original')
assert all(hashlib.sha256((root/m['name']).read_bytes()).hexdigest()==m['sha256'] for m in manifest)
result={'tasks':len(tasks),'completed_planning':3,'development_progress':0,'all_required_screens_precede_production':True,'formula_errors':errors,'source_files_unchanged':len(manifest),'percent_by_count':3/127,'open_decision_records':52,'validation':'Structure, dependencies, original hashes, saved XLSX formulas/caches, plus Artifact Tool input-change and render checks. Native Excel app not exercised.'}
Path('work/verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
