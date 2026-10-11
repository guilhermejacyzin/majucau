"""Verifica somente os artefatos incluídos, sem consultar sistemas ou fontes externas."""
import ast
import hashlib
import json
from collections import Counter
import re
from pathlib import Path
from urllib.parse import unquote

import openpyxl

root = Path(__file__).resolve().parents[1]
repo = root.parents[1]
tasks = json.loads((root / 'work/backlog_v2.json').read_text(encoding='utf-8'))
decisions = json.loads((root / 'work/decisions_v2.json').read_text(encoding='utf-8'))
by = {t['key']:t for t in tasks}
ids = [t['id'] for t in tasks]
assert ids == [f'BK-{i:03}' for i in range(127)]

def ancestors(key, stack=()):
    assert key not in stack, ('cycle', stack, key)
    result = set()
    for dep in by[key]['deps']:
        assert dep in by, dep
        result.add(dep)
        result.update(ancestors(dep, stack + (key,)))
    return result

for key in by:
    ancestors(key)
required = ['losses','supplementaryingest','usersscreen','parametersscreen',
    'integrationsscreen','reconciliationscreen','receivablesscreen','payablesscreen',
    'cashscreen','drescreen','trialbalancescreen','calculatorscreen','applicationsscreen',
    'pricing','logisticsanalysis','inventoryscreen','productionscreen','purchasesscreen',
    'projectionsscreen','projectionexport','pendingscreen','alertsscreen','executivescreen',
    'extrascreens','approvedexports','cyclemetrics','restoreexercise']
assert set(required) <= ancestors('readiness')
assert 'cocoasignals' not in ancestors('purchasesscreen')
assert 'inputprojection' in ancestors('purchaseeligibility')
assert 'inventoryprojection' not in ancestors('productionprojection')
assert all(t['goal'] and t['action'] and t['accept'] and t['impact'] and len(t['steps']) >= 3 for t in tasks)
assert all(0 <= t['progress'] <= 1 for t in tasks)
assert all(t['progress'] == 1 and t['evidence'] for t in tasks if t['status'] == 'Concluída')
assert not any(t['status'] == 'Aguarda decisão' for t in tasks)
assert len(decisions) == 108
md = (root / 'outputs/MAJUCAU_BACKLOG_DETALHADO_v2.md').read_text(encoding='utf-8')
headers = list(re.finditer(r'^### (BK-\d+) — ([^\n]+)', md, re.M))
assert [match.group(1) for match in headers] == ids
markdown_tasks = {}
for index, match in enumerate(headers):
    end = headers[index + 1].start() if index + 1 < len(headers) else len(md)
    body = md[match.end():end]
    status = re.search(r'\*\*Situação:\*\*\s*(Concluída|Em andamento|Não iniciada)\.', body)
    progress = re.search(r'\*\*Evolução:\*\*\s*(\d+(?:[.,]\d+)?)%', body)
    assert status and progress, match.group(1)
    markdown_tasks[match.group(1)] = (
        status.group(1),
        float(progress.group(1).replace(',', '.')) / 100,
    )
for task in tasks:
    status, progress = markdown_tasks[task['id']]
    assert task['status'] == status, (task['id'], task['status'], status)
    assert abs(task['progress'] - progress) < 0.000001, (task['id'], task['progress'], progress)
status_counts = Counter(task['status'] for task in tasks)
assert set(status_counts) <= {'Concluída', 'Em andamento', 'Não iniciada'}
assert sum(status_counts.values()) == len(tasks)
estimated_progress = sum(task['progress'] * task['weight'] for task in tasks) / sum(task['weight'] for task in tasks)
assert 'D+1' not in md
planning_accepted = sum(t['weight'] for t in tasks if t['kind'] == 'Planejamento' and t['status'] == 'Concluída' and t['evidence'])

for xlsx in (root / 'outputs').glob('*.xlsx'):
    cached = openpyxl.load_workbook(xlsx, data_only=True)
    assert not [(s.title,c.coordinate,c.value) for s in cached for row in s for c in row if c.data_type == 'e']
    if xlsx.name.endswith('_v2.xlsx'):
        formula = openpyxl.load_workbook(xlsx, data_only=False)
        sheet = cached['Evolucao']
        development_weight = sum(t['weight'] for t in tasks if t['kind'] == 'Desenvolvimento')
        development_accepted = sum(t['weight'] for t in tasks if t['kind'] == 'Desenvolvimento' and t['status'] == 'Concluída' and t['evidence'])
        accepted_count = sum(t['status'] == 'Concluída' and bool(t['evidence']) for t in tasks)
        assert [sheet[f'B{n}'].value for n in [6,7,10,11]] == [127,planning_accepted,10,accepted_count]
        assert abs(sheet['B8'].value - estimated_progress) < 1e-10
        assert abs(sheet['B9'].value - development_accepted/development_weight) < 1e-10
        assert all(formula['Tarefas'].cell(i+6,13).data_type == 'f' for i in range(127))
        assert all(formula['Tarefas'].cell(i+6,14).data_type == 'f' for i in range(127))
        assert [formula['Tarefas'].cell(i+6,7).value for i in range(127)] == [t['progress'] for t in tasks]
        assert [formula['Tarefas'].cell(i+6,1).value for i in range(127)] == ids
        assert [formula['Tarefas'].cell(i+6,3).value for i in range(127)] == [t['title'] for t in tasks]

links = 0
for path in repo.rglob('*.md'):
    if '.git' in path.parts or 'artifacts' in path.parts or 'node_modules' in path.parts:
        continue
    text = path.read_text(encoding='utf-8')
    # O README anterior do proprietário conserva referências não incluídas nesta entrega.
    # Aqui conferimos somente o bloco documental acrescentado por este pacote.
    if path == repo / 'README.md' and '<!-- MAJUCAU_PLANEJAMENTO_CONVERSA -->' in text:
        text = text.split('<!-- MAJUCAU_PLANEJAMENTO_CONVERSA -->',1)[1]
    assert not re.search(r'C:[/\\]Users[/\\]', text), str(path)
    for target in re.findall(r'\]\(<?([^\n)>]+)>?\)', text):
        if re.match(r'^[a-z][a-z0-9+.-]*:', target, re.I) or target.startswith('#'):
            continue
        dest = (path.parent / unquote(target.split('#')[0])).resolve()
        assert dest.is_relative_to(repo), (path, target)
        assert dest.exists(), (path, target)
        links += 1
for path in root.glob('work/*.py'):
    ast.parse(path.read_text(encoding='utf-8'), filename=str(path))

manifest = root / 'MANIFESTO.json'
if manifest.exists():
    for entry in json.loads(manifest.read_text(encoding='utf-8'))['files']:
        p = repo / entry['path']
        assert p.is_file(), entry['path']
        assert hashlib.sha256(p.read_bytes()).hexdigest() == entry['sha256'], entry['path']

print(json.dumps({'tasks':127,'decisions':108,'technical_fronts':10,
    'completed_planning':planning_accepted,
    'completed_with_evidence':sum(t['status'] == 'Concluída' and bool(t['evidence']) for t in tasks),
    'estimated_progress_percent':round(estimated_progress*100,1),'acyclic_dependencies':True,
    'formula_errors':0,'valid_local_markdown_links':links,
    'original_sources_checked':False,'product_tests_executed':False}, ensure_ascii=False))
