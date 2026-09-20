import json,re,hashlib
from pathlib import Path
import openpyxl
W=Path('work');O=Path('outputs')
t=json.loads((W/'backlog_v2.json').read_text(encoding='utf-8'));by={x['key']:x for x in t}
d=json.loads((W/'decisions_v2.json').read_text(encoding='utf-8'))
ids=[x['id'] for x in t]
assert ids==[f'BK-{i:03}' for i in range(127)]
def ancestors(k,stack=()):
    assert k not in stack,('cycle',stack,k)
    s=set()
    for v in by[k]['deps']:
        assert v in by
        s.add(v);s.update(ancestors(v,stack+(k,)))
    return s
for k in by:ancestors(k)
required=['losses','supplementaryingest','usersscreen','parametersscreen','integrationsscreen','reconciliationscreen','receivablesscreen','payablesscreen','cashscreen','drescreen','trialbalancescreen','calculatorscreen','applicationsscreen','pricing','logisticsanalysis','inventoryscreen','productionscreen','purchasesscreen','projectionsscreen','projectionexport','pendingscreen','alertsscreen','executivescreen','extrascreens','approvedexports','cyclemetrics','restoreexercise']
assert set(required)<=ancestors('readiness'),set(required)-ancestors('readiness')
assert 'cocoasignals' not in ancestors('purchasesscreen')
assert 'inputprojection' in ancestors('purchaseeligibility')
assert 'inventoryprojection' not in ancestors('productionprojection')
assert not set(['losses','receivableingest','historicalcost','logisticsingest'])&set(by['dreengine']['deps'])
assert sum(x['progress'] for x in t)==3
assert all(x['progress']==0 for x in t if x['kind']=='Desenvolvimento')
assert not any(x['status']=='Aguarda decisão' for x in t)
assert all(x['goal'] and x['action'] and x['accept'] and x['impact'] and len(x['steps'])>=3 for x in t)
md=(O/'MAJUCAU_BACKLOG_DETALHADO_v2.md').read_text(encoding='utf-8')
assert re.findall(r'^### (BK-\d+) —',md,re.M)==ids
assert 'D+1' not in md
p=O/'MAJUCAU_BACKLOG_E_EVOLUCAO_v2.xlsx'
f=openpyxl.load_workbook(p,data_only=False);c=openpyxl.load_workbook(p,data_only=True)
assert c['Evolucao']['B6'].value==127
assert c['Evolucao']['B7'].value==3
assert abs(c['Evolucao']['B8'].value-3/127)<1e-10
assert c['Evolucao']['B9'].value==0
assert c['Evolucao']['B10'].value==10
assert all(f['Tarefas'].cell(i+6,7).data_type=='f' for i in range(127))
assert [f['Tarefas'].cell(i+6,1).value for i in range(127)]==ids
assert len(d)==108
assert not [(s.title,x.coordinate,x.value) for s in c for row in s for x in row if x.data_type=='e']
for x in json.loads((W/'source_extract/manifest.json').read_text(encoding='utf-8')):
    assert hashlib.sha256((Path(r'sources/private/material-original')/x['name']).read_bytes()).hexdigest()==x['sha256']
for v in [1,2]:
    m=json.loads((W/f'revision_v{v}.json').read_text(encoding='utf-8'))
    assert hashlib.sha256(Path(m['source']).read_bytes()).hexdigest()==m['sha256']
result={'tasks':127,'revised_tasks':70,'technical_fronts':10,'business_questions_reopened':0,'completed_planning':3,'development_progress':0,'acyclic_dependencies':True,'losses_required_before_release':True,'source_files_unchanged':19,'formula_errors':0,'validation':'Planning artifacts only; no product implementation or live integration tests.'}
(W/'verification_v2.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result))
